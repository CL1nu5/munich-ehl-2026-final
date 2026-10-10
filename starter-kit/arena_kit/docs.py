"""Read the documents suppliers attach: PDF and XLSX to text, a quote PDF to structured terms."""

import io
import re

LABELS = {
    "Unit price (EUR)": "unit_price",
    "Payment terms (days net)": "payment_days",
    "Lead time (days)": "lead_time_days",
    "Volume per year (pcs)": "volume_per_year",
    "Contract term (months)": "contract_months",
    "Price valid from": "effective_date",
    "Index clause": "index_clause",
    "Incoterm": "incoterm",
    "Tooling (EUR)": "tooling_eur",
    "Quality commitment (ppm)": "ppm_commitment",
    "Chargeback per complaint (EUR)": "chargeback_eur_per_complaint",
}
SECTIONS = {
    "Offered terms": "stated",
    "Our standard terms": "standard",
    "Terms fixed by your enquiry": "fixed",
    "Agreed terms": "agreed",
}
CLAUSE = re.compile(r"(\S+), (\d+)% material, deadband ([\d.]+) %, base (\S+)")


def read_document(name, data):
    """The text of an attachment: PDF pages, XLSX rows (formulas as written), or plain text."""
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data))
        return "\n".join(
            f"[{ws.title}]\n"
            + "\n".join(" | ".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True))
            for ws in wb.worksheets
        )
    return data.decode("utf-8", "replace")


def _value(field, text):
    if field == "index_clause":
        m = CLAUSE.match(text)
        if not m:
            return None
        series, share, deadband, base = m.groups()
        return {
            "series": series,
            "material_share": int(share) / 100,
            "deadband_pct": float(deadband),
            "base_month": base,
        }
    if field in ("effective_date", "incoterm"):
        return text
    return float(text) if "eur" in field or field == "unit_price" else int(float(text))


def quote_terms(text):
    """Terms per section of a quote or order confirmation PDF: stated, standard, fixed, agreed."""
    out, section = {}, None
    for line in text.splitlines():
        line = line.strip()
        section = next((key for head, key in SECTIONS.items() if line.startswith(head)), section)
        label = next((lab for lab in LABELS if line.startswith(lab + " ")), None)
        if section and label:
            out.setdefault(section, {})[LABELS[label]] = _value(LABELS[label], line[len(label) :].strip())
    return out


def standard_terms(events):
    """The supplier's standard terms from the latest quote PDF in a thread (empty if none)."""
    for msg in reversed(events):
        for att in msg.get("attachments") or []:
            if att.get("name", "").startswith("quote") and att.get("text"):
                return quote_terms(att["text"]).get("standard", {})
    return {}
